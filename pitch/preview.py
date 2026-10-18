"""Render an approximate HTML preview of the deck (with the Barlow font) for layout checks.

Run: uv run --with python-pptx python pitch/preview.py
Output: pitch/out/preview/slide-N.html (screenshot them with headless Chrome).
"""

import re
from pathlib import Path

from pptx import Presentation
from pptx.enum.dml import MSO_FILL
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

ROOT = Path(__file__).parent
OUT = ROOT / "out" / "preview"
FONTS = ROOT.parent / "web" / "node_modules" / "@fontsource" / "barlow" / "files"
PX = 100 / 914400  # 100 px per inch


def color(c) -> str:
    return f"#{c}"


def alpha_of(el) -> float:
    a = el.find(".//" + qn("a:alpha"))
    return int(a.get("val")) / 100000 if a is not None else 1.0


def fill_css(shape) -> str:
    try:
        if shape.fill.type == MSO_FILL.SOLID:
            rgb = shape.fill.fore_color.rgb
            a = alpha_of(shape.fill._xPr.find(qn("a:solidFill")))
            r, g, b = int(str(rgb)[0:2], 16), int(str(rgb)[2:4], 16), int(str(rgb)[4:6], 16)
            return f"background: rgba({r},{g},{b},{a});"
    except Exception:
        pass
    return ""


def line_css(shape) -> str:
    try:
        ln = shape.line._ln
        if ln is None or ln.find(qn("a:noFill")) is not None:
            return ""
        rgb = shape.line.color.rgb
        w = (shape.line.width or 12700) * PX * 0.75 / 0.75
        dash = "dashed" if ln.find(qn("a:prstDash")) is not None else "solid"
        return f"border: {max(1, w):.1f}px {dash} #{rgb};"
    except Exception:
        return ""


def text_html(tf) -> str:
    out = []
    for p in tf.paragraphs:
        align = {None: "left", 1: "left", 2: "center", 3: "right"}.get(p.alignment if p.alignment is None else int(p.alignment), "left")
        lh = p.line_spacing if isinstance(p.line_spacing, float) else 1.1
        sa = p.space_after.pt * 100 / 72 if p.space_after is not None else 0
        spans = []
        for r in p.runs:
            f = r.font
            size = f.size.pt * 100 / 72 if f.size else 24
            col = f"#{f.color.rgb}" if f.color and f.color.type is not None else "#E5E7E1"
            weight = 600 if f.bold else 400
            spans.append(f'<span style="font-size:{size:.1f}px;color:{col};font-weight:{weight}">{r.text.replace("&", "&amp;").replace("<", "&lt;")}</span>')
        out.append(f'<p style="margin:0 0 {sa:.1f}px;text-align:{align};line-height:{lh * 1.2:.2f}">{"".join(spans) or "&nbsp;"}</p>')
    return "".join(out)


def main():
    prs = Presentation(ROOT / "out" / "WildfireTwinLA-pitch.pptx")
    OUT.mkdir(parents=True, exist_ok=True)
    W, H = prs.slide_width * PX, prs.slide_height * PX
    for n, slide in enumerate(prs.slides, 1):
        parts = []
        for k, sh in enumerate(slide.shapes):
            x, y = (sh.left or 0) * PX, (sh.top or 0) * PX
            w, h = (sh.width or 0) * PX, (sh.height or 0) * PX
            pos = f"position:absolute;left:{x:.1f}px;top:{y:.1f}px;width:{w:.1f}px;height:{h:.1f}px;"
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                name = f"s{n}-{k}.{sh.image.ext}"
                (OUT / name).write_bytes(sh.image.blob)
                parts.append(f'<img src="{name}" style="{pos}{line_css(sh)}box-sizing:border-box">')
            elif sh.shape_type == MSO_SHAPE_TYPE.TABLE or getattr(sh, "has_table", False):
                t = sh.table
                cols = [c.width * PX for c in t.columns]
                rows = []
                for row in t.rows:
                    cells = "".join(f'<td style="width:{cols[i]:.0f}px;padding:6px 12px;vertical-align:top;{fill_css(c)}">{text_html(c.text_frame)}</td>' for i, c in enumerate(row.cells))
                    rows.append(f"<tr>{cells}</tr>")
                parts.append(f'<table style="{pos}height:auto;border-collapse:separate;border-spacing:2px">{"".join(rows)}</table>')
            elif sh.shape_type == MSO_SHAPE_TYPE.LINE or sh.__class__.__name__ == "Connector":
                x0, y0, x1, y1 = sh.begin_x * PX, sh.begin_y * PX, sh.end_x * PX, sh.end_y * PX
                rgb = sh.line.color.rgb
                dash = 'stroke-dasharray="6 5"' if sh.line._ln.find(qn("a:prstDash")) is not None else ""
                parts.append(
                    f'<svg style="position:absolute;left:0;top:0;overflow:visible" width="{W}" height="{H}"><defs><marker id="m{k}" markerWidth="8" markerHeight="8" refX="7" refY="4" orient="auto"><path d="M0,0 L8,4 L0,8 z" fill="#{rgb}"/></marker></defs>'
                    f'<line x1="{x0:.1f}" y1="{y0:.1f}" x2="{x1:.1f}" y2="{y1:.1f}" stroke="#{rgb}" stroke-width="2" {dash} marker-end="url(#m{k})"/></svg>'
                )
            else:
                radius = "border-radius:10px;" if "roundRect" in sh._element.xml[:2000] else ""
                inner = text_html(sh.text_frame) if sh.has_text_frame and sh.text_frame.text else ""
                parts.append(f'<div style="{pos}{fill_css(sh)}{line_css(sh)}{radius}box-sizing:border-box;overflow:visible">{inner}</div>')
        bg = "#1E1A17"
        try:
            bg = f"#{slide.background.fill.fore_color.rgb}"
        except Exception:
            pass
        fonts = "".join(
            f'@font-face{{font-family:Barlow;font-weight:{w};src:url("{(FONTS / f"barlow-latin-{w}-normal.woff2").as_uri()}")}}' for w in (400, 500, 600)
        )
        html = f'<!doctype html><meta charset="utf-8"><style>{fonts}body{{margin:0;background:#000}}#s{{position:relative;width:{W:.0f}px;height:{H:.0f}px;background:{bg};overflow:hidden;font-family:Barlow,sans-serif}}</style><div id="s">{"".join(parts)}</div>'
        (OUT / f"slide-{n}.html").write_text(html)
    print("wrote", len(prs.slides), "slides to", OUT)


if __name__ == "__main__":
    main()
