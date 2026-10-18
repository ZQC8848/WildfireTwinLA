"""Build the 10-slide WildfireTwinLA pitch deck (.pptx) from screenshots and facts.json.

Run: uv run --with python-pptx --with segno python pitch/build_deck.py
Output: pitch/out/WildfireTwinLA-pitch.pptx (import into Google Slides by uploading to Drive).
"""

import json
from pathlib import Path

import segno
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).parent
APP = ROOT / "assets" / "app"
WEB = ROOT / "assets" / "web"
OUT = ROOT / "out"
FACTS = json.loads((ROOT / "facts.json").read_text())
NOTES = json.loads((ROOT / "notes.json").read_text())

FONT = "Barlow"
BG = RGBColor(0x1E, 0x1A, 0x17)
PANEL = RGBColor(0x2B, 0x23, 0x1F)
INK = RGBColor(0xE5, 0xE7, 0xE1)
MUTED = RGBColor(0xA9, 0xAB, 0xA4)
EMBER = RGBColor(0xF2, 0x6A, 0x1B)
MARINE = RGBColor(0x6F, 0x9D, 0xBF)
PINK = RGBColor(0xD6, 0x3A, 0x73)
AMBER = RGBColor(0xE8, 0xA3, 0x3A)
GRAY = RGBColor(0x8C, 0x8E, 0x88)

RESPONSIBLE = "Experimental, unofficial model outputs. Demo contacts and call outcomes are fictional. Follow official evacuation instructions."

CHIPS = {
    "OBSERVED": MARINE,
    "MODELLED": EMBER,
    "REPLAY": GRAY,
    "SIMULATED CONTACT": PINK,
    "PLANNED": INK,
    "PROPOSED": INK,
    "ILLUSTRATIVE": GRAY,
    "EXTERNAL SOURCE": MARINE,
}

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def set_alpha(shape, alpha: float):
    """Fill transparency (0 = transparent, 1 = opaque)."""
    fill = shape.fill._xPr.find(qn("a:solidFill"))
    color = fill[0]
    for old in color.findall(qn("a:alpha")):
        color.remove(old)
    a = etree.SubElement(color, qn("a:alpha"))
    a.set("val", str(int(alpha * 100000)))


def rect(slide, x, y, w, h, fill=None, line=None, dash=False, radius=False, line_w=1.25):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        shape.adjustments[0] = 0.12
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_w)
        if dash:
            ln = shape.line._get_or_add_ln()
            d = etree.SubElement(ln, qn("a:prstDash"))
            d.set("val", "dash")
    shape.shadow.inherit = False
    return shape


def text(slide, x, y, w, h, runs, size=24, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.1):
    """runs: str, or list of paragraphs; each paragraph is str or list of (text, {opts})."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    paras = [runs] if isinstance(runs, str) else runs
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        parts = [(para, {})] if isinstance(para, str) else para
        for t, o in parts:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = FONT
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", False)
            f.color.rgb = o.get("color", color)
        if isinstance(para, list) and para and para[0][1].get("space_after"):
            p.space_after = Pt(para[0][1]["space_after"])
    return box


def bullets(slide, x, y, w, h, items, size=24, gap=14, marker_color=PINK):
    paras = [[("▪  ", {"color": marker_color, "size": size}), (t, {"size": size, "space_after": gap})] for t in items]
    for para in paras:
        para[0][1]["space_after"] = gap
    return text(slide, x, y, w, h, paras, size=size)


def chip(slide, x, y, label, sub=None):
    color = CHIPS.get(label, INK)
    w = 0.16 + 0.105 * len(label)
    dashed = label in ("PLANNED", "PROPOSED")
    rect(slide, x, y, w, 0.32, fill=None, line=color, dash=dashed, radius=True, line_w=1.25)
    text(slide, x, y + 0.035, w, 0.27, [[(label, {"bold": True, "size": 12, "color": color})]], align=PP_ALIGN.CENTER)
    if sub:
        text(slide, x + w + 0.12, y + 0.035, 6, 0.3, [[(sub, {"size": 13, "color": MUTED})]])
    return w


def picture(slide, path, x, y, w=None, h=None, border=True):
    if not Path(path).exists():
        # 截图尚未生成时放一个占位框，便于先检查版式
        box = rect(slide, x, y, w or 4, h or 2.5, fill=PANEL, line=AMBER, dash=True)
        text(slide, x, y + 0.2, w or 4, 0.4, [[("PENDING: " + Path(path).name, {"size": 12, "color": AMBER})]], align=PP_ALIGN.CENTER)
        return box
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)
    if border:
        pic.line.color.rgb = RGBColor(0x45, 0x3C, 0x37)
        pic.line.width = Pt(0.75)
    return pic


def arrow(slide, x0, y0, x1, y1, color=MUTED, dash=False, width=2.0):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x0), Inches(y0), Inches(x1), Inches(y1))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dash:
        d = etree.SubElement(ln, qn("a:prstDash"))
        d.set("val", "dash")
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    return c


def base(n, title, subtitle=None, footer=True):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    text(s, 0.6, 0.38, 12.1, 0.8, [[(title, {"bold": True, "size": 38})]])
    if subtitle:
        text(s, 0.6, 1.08, 12.1, 0.5, [[(subtitle, {"size": 20, "color": MUTED})]])
    # 页脚：编号与负责任使用说明
    text(s, 12.2, 7.06, 0.6, 0.3, [[(str(n), {"size": 12, "color": MUTED})]], align=PP_ALIGN.RIGHT)
    if footer:
        text(s, 0.6, 7.06, 11.2, 0.3, [[(RESPONSIBLE, {"size": 12, "color": MUTED})]])
    s.notes_slide.notes_text_frame.text = NOTES[str(n)]
    return s


F = FACTS

# ---------- 1. Why calls matter ----------
s = base(1, "A warning sent isn’t a person reached.", "WildfireTwinLA — from fire information to prioritized outreach.", footer=False)
bg = picture(s, APP / "P2a-replay-overview.png", 0, 0, w=13.333, border=False)
s.shapes._spTree.remove(bg._element)
s.shapes._spTree.insert(2, bg._element)
veil = rect(s, 0, 0, 13.333, 7.5, fill=BG)
set_alpha(veil, 0.78)
s.shapes._spTree.remove(veil._element)
s.shapes._spTree.insert(3, veil._element)
bullets(s, 0.6, 2.0, 6.2, 3.2, ["Confirm that someone received the message.", "Let them ask for repetition or help.", "Keep unanswered contacts visible."], size=28, gap=22)
labels = [("Acknowledged", "P1-confirmed.png"), ("No answer", "P1-missed.png"), ("Needs a person", "P1-help.png")]
y = 1.95
for label, f in labels:
    text(s, 7.4, y, 5.2, 0.35, [[(label.upper(), {"bold": True, "size": 14, "color": INK})]])
    pic = picture(s, APP / f, 7.4, y + 0.38, w=5.3)
    y += 0.38 + pic.height / 914400 + 0.25
chip(s, 7.4, y + 0.02, "SIMULATED CONTACT", "Outcomes from the demo dispatcher; fictional residents")
text(s, 0.6, 7.06, 12, 0.3, [[("Demo contacts and call outcomes are fictional.", {"size": 12, "color": MUTED})]])

# ---------- 2. Real fires ----------
s = base(2, "Real fires. A local decision.", "Public fire observations, geographic context and experimental spread scenarios.")
picture(s, APP / "P2a-replay-overview.png", 0.6, 1.7, w=8.1)
chip(s, 0.6, 6.35, "REPLAY", None)
text(
    s,
    1.55,
    6.36,
    7.2,
    0.6,
    [[(F["replay_caption"], {"size": 12, "color": MUTED})]],
)
rect(s, 9.0, 1.7, 3.75, 5.15, fill=PANEL, radius=True)
text(s, 9.2, 1.82, 3.4, 0.4, [[("Regional context — CAL FIRE", {"bold": True, "size": 17})]])
y = 2.28
for inc in F["calfire"]:
    text(s, 9.2, y, 3.4, 0.9, [[(inc["name"], {"bold": True, "size": 17, "color": EMBER if inc["name"].startswith("Norte") else INK})], [(inc["where"], {"size": 13, "color": MUTED})], [(inc["stats"], {"size": 15})]], spacing=1.0)
    y += 0.86
text(s, 9.2, y - 0.02, 3.4, 0.7, [[(F["calfire_note"], {"size": 11, "color": MUTED})]], spacing=1.0)
if F.get("news"):
    picture(s, WEB / F["news"]["file"], 9.2, 5.52, w=3.35)

# ---------- 3. Who comes first ----------
s = base(3, "Limited staff. Many households. Who comes first?", footer=False)
bullets(s, 0.6, 1.5, 6.7, 2.6, ["User: the overnight call-line supervisor.", "Decision: prioritize outbound follow-up.", "Delivery: simulated today; Retell integration planned."], size=22, gap=12)
# 无序列表 → 有序队列 → 计划中的投递
rect(s, 0.6, 4.25, 2.6, 2.45, fill=PANEL, radius=True)
text(s, 0.75, 4.33, 2.4, 0.35, [[("Households near the fire", {"bold": True, "size": 14})]])
for i, wdt in enumerate([1.9, 1.4, 2.1, 1.6, 1.8]):
    rect(s, 0.8, 4.78 + i * 0.33, wdt, 0.18, fill=RGBColor(0x55, 0x4C, 0x47), radius=True)
chip(s, 0.75, 6.32, "ILLUSTRATIVE")
arrow(s, 3.35, 5.45, 4.05, 5.45, color=INK)
picture(s, APP / "P3-queue-top5.png", 4.2, 3.95, h=2.75)
text(s, 4.2, 6.75, 3.6, 0.3, [[("Ordered by estimated fire arrival", {"size": 13, "color": MUTED})]])
arrow(s, 7.65, 5.3, 8.9, 5.3, color=INK, dash=True)
rect(s, 9.05, 4.55, 3.6, 1.5, fill=None, line=INK, dash=True, radius=True)
text(s, 9.2, 4.7, 3.3, 1.2, [[("Retell voice call", {"bold": True, "size": 20})], [("Approved update; outcome returned to the queue", {"size": 14, "color": MUTED})]])
chip(s, 9.05, 6.2, "PLANNED", "Calls are simulated today")
picture(s, APP / "P3-progress.png", 7.45, 1.5, w=5.2)
text(s, 7.45, 3.78, 5.2, 0.3, [[("Live counts from the simulated dispatcher", {"size": 12, "color": MUTED})]])
text(s, 0.6, 7.06, 12, 0.3, [[("Ranking reflects estimated fire arrival, not conversational urgency or rescue priority. Demo contacts are fictional.", {"size": 12, "color": MUTED})]])

# ---------- 4. Queue logic ----------
s = base(4, "Nearby homes become a clear outreach queue.")
steps = ["Home location", "Distance + wind (or model ETA)", "Estimated arrival", "Sorted queue", "Call status"]
x = 0.6
for i, st in enumerate(steps):
    w = 0.34 + 0.1 * len(st)
    rect(s, x, 1.35, w, 0.5, fill=PANEL, line=PINK if i >= 3 else EMBER, radius=True)
    text(s, x, 1.44, w, 0.4, [[(st, {"size": 15, "bold": True})]], align=PP_ALIGN.CENTER)
    x += w
    if i < len(steps) - 1:
        arrow(s, x + 0.04, 1.6, x + 0.3, 1.6, color=MUTED)
        x += 0.34
picture(s, APP / "P4-calling-closeup.png", 0.6, 2.1, h=3.85)
text(s, 0.6, 6.0, 5.4, 0.35, [[("Each home’s status sits above its roof (ring = queued).", {"size": 13, "color": MUTED})]], spacing=1.0)
picture(s, APP / "P3-queue-top5.png", 6.55, 2.1, h=3.85)
text(s, 6.55, 6.0, 4.8, 0.35, [[("Rank · masked address · resident · flags · arrival", {"size": 13, "color": MUTED})]], spacing=1.0)
text(s, 11.35, 2.15, 1.45, 3.6, [[("Top of the queue", {"bold": True, "size": 15})], [("earliest expected fire arrival", {"size": 13, "color": MUTED, "space_after": 10})], [("“Retry”", {"bold": True, "size": 15})], [("no answer on the first call", {"size": 13, "color": MUTED, "space_after": 10})], [("Farther away is not the same as safe.", {"size": 13, "color": AMBER})]], spacing=1.0)
text(
    s,
    0.6,
    6.45,
    12.1,
    0.55,
    [[("Ranked by estimated fire arrival at each home (nearest simulated burning cell; simulation seeded from reported ignitions and VIIRS hotspots). Access-needs flag moves a home up 5 min. Not a complete danger assessment.", {"size": 12, "color": INK})]],
    spacing=1.0,
)

# ---------- 5. Preparation ----------
s = base(5, "Before the emergency: where should we prepare?", footer=False)
bullets(s, 0.6, 1.55, 4.6, 4.0, ["Intended user: community planner or Fire Safe Council.", "Explore preventive treatment locations.", "Compare scenarios before allocating resources."], size=24, gap=18)
picture(s, APP / "P6-plan-break.png", 5.35, 1.5, w=7.4)
chip(s, 5.35, 5.82, "MODELLED", "Plan mode: a fuel break drawn downwind of the Norte ignition point")
text(s, 0.6, 7.06, 12, 0.3, [[("Prevention planning, not a guarantee. Prospective users are not confirmed partners.", {"size": 12, "color": MUTED})]])

# ---------- 6. Prevention comparison ----------
s = base(6, "Test a treatment. Compare the simulated outcome.", "Fuel break · Prescribed burn · Thinning · Home hardening")
picture(s, APP / "P7-full.png", 0.6, 1.6, w=7.5)
picture(s, APP / "P7c-compare-table.png", 8.4, 1.6, w=4.35)
text(s, 0.6, 5.88, 12.1, 0.6, [[(F["compare_takeaway"], {"size": 19, "bold": True})]], spacing=1.0)
chip(s, 0.6, 6.62, "MODELLED", F["compare_caption"])

# ---------- 7. Data ----------
s = base(7, "What powers today’s prototype?", footer=False)
rows = [
    ("Input", "Function", "Source / status"),
    ("Terrain, fuels, buildings", "Geographic context and home locations", "USGS 3DEP, LANDFIRE FBFM40, OpenStreetMap (50,408 buildings), Sentinel-2 imagery (Oct 14, 2026) — committed real data"),
    ("Satellite fire observations", "Reported hotspot locations and times", "NASA FIRMS VIIRS — replay of 4 passes on Oct 16; live FIRMS/GOES pollers built, not configured"),
    ("Wind", "Replay conditions and simulation input", "Archived NOAA HRRR (5×5 points, hourly); WindNinja terrain-wind library"),
    ("Resident contacts", "Demonstrate the outreach workflow", "Synthetic, one per building; fictional (949) 555-01xx numbers; never dialled"),
]
tbl = s.shapes.add_table(len(rows), 3, Inches(0.6), Inches(1.5), Inches(12.1), Inches(4.4)).table
tbl.columns[0].width = Inches(3.0)
tbl.columns[1].width = Inches(3.5)
tbl.columns[2].width = Inches(5.6)
for r, row in enumerate(rows):
    for c, val in enumerate(row):
        cell = tbl.cell(r, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = PANEL if r else RGBColor(0x3A, 0x31, 0x2C)
        cell.margin_left = cell.margin_right = Inches(0.12)
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = val
        run.font.name = FONT
        run.font.size = Pt(15 if r and c == 2 else 18)
        run.font.bold = r == 0 or c == 0
        run.font.color.rgb = INK
text(s, 0.6, 6.15, 12.1, 0.5, [[("Observed data, model outputs and fictional contacts are labelled separately.", {"size": 24, "bold": True})]])
text(s, 0.6, 6.75, 12.1, 0.55, [[("Community areas are approximate, not official evacuation zones. Replay evacuation orders: OCFA, as reported by CBS Los Angeles and KTLA (times approximate). Sentinel-2 is periodic imagery, not real-time monitoring.", {"size": 12, "color": MUTED})]], spacing=1.0)

# ---------- 8. Deployment ----------
s = base(8, "What would an operational deployment need?", footer=False)
bullets(s, 0.6, 1.45, 6.3, 2.2, ["Verified, authorized contact records.", "Agency-approved zones and message sources.", "Additional observations and independent validation."], size=22, gap=10)
cols = [("Residents", "Opt-in enrollment, language, consent"), ("Agency partners", "Approved zones, messages, contact data"), ("Field observations", "Validation of fire position and model")]
for i, (h, sub) in enumerate(cols):
    x = 7.1 + (i % 3) * 1.9
    rect(s, x, 1.45, 1.78, 1.25, fill=PANEL, radius=True)
    text(s, x + 0.1, 1.52, 1.6, 1.1, [[(h, {"bold": True, "size": 14})], [(sub, {"size": 11, "color": MUTED})]], spacing=1.0)
    arrow(s, x + 0.89, 2.75, 9.9, 3.25, color=MUTED, dash=True, width=1.5)
rect(s, 8.65, 3.3, 2.5, 0.5, fill=None, line=PINK, radius=True)
text(s, 8.65, 3.38, 2.5, 0.4, [[("WildfireTwinLA", {"bold": True, "size": 15})]], align=PP_ALIGN.CENTER)
chip(s, 11.3, 3.38, "PROPOSED")
# Twist 回应
for i, tw in enumerate([F["twist1"], F["twist2"]]):
    x = 0.6 + i * 6.15
    rect(s, x, 4.0, 6.0, 2.95, fill=PANEL, line=AMBER, radius=True)
    paras = [[(tw["title"], {"bold": True, "size": 17, "color": AMBER, "space_after": 6})]]
    paras += [[(ln, {"size": 14, "space_after": 6})] for ln in tw["lines"]]
    paras += [[(tw["source"], {"size": 10, "color": MUTED})]]
    text(s, x + 0.18, 4.1, 5.65, 2.8, paras, spacing=1.0)
text(s, 0.6, 7.06, 12, 0.3, [[("No agency relationship or funding award is implied. Building footprints do not establish a resident’s phone number or permission to call.", {"size": 12, "color": MUTED})]])

# ---------- 9. No answer ----------
s = base(9, "No answer does not mean safe.")
flow = [
    ("Call attempted", None),
    ("No answer → retry once after 3 min", "SIMULATED CONTACT"),
    ("Still no answer → text moved up the queue", "SIMULATED CONTACT"),
    ("Text delivered ≠ acknowledged", None),
    ("Reply captured?  Yes → record · No → unresolved", "PLANNED"),
]
y = 1.4
for i, (label, tag) in enumerate(flow):
    unresolved = i == len(flow) - 1
    rect(s, 0.6, y, 6.6, 0.62, fill=PANEL, line=AMBER if unresolved else None, dash=unresolved, radius=True)
    text(s, 0.78, y + 0.12, 6.3, 0.45, [[(label, {"size": 18, "bold": i in (0, 4)})]])
    if tag:
        chip(s, 7.35, y + 0.15, tag)
    if i < len(flow) - 1:
        arrow(s, 1.2, y + 0.64, 1.2, y + 0.88, color=MUTED)
    y += 0.9
picture(s, APP / "P1-missed.png", 9.5, 1.4, w=3.25)
picture(s, APP / "P10-retry-row.png", 9.5, 2.25, w=3.25)
text(s, 9.5, 2.9, 3.25, 0.9, [[("Line card and queue row in the demo: the retry stays visible to the operator.", {"size": 13, "color": AMBER})]], spacing=1.0)
strip = [("Stale fire data → show age", "PLANNED"), ("Missing observations → simulation remains an estimate", None), ("Phone/SMS outage → delivery unavailable", None)]
x = 0.6
for label, tag in strip:
    rect(s, x, 5.95, 3.95, 0.85, fill=None, line=RGBColor(0x55, 0x4C, 0x47), radius=True)
    text(s, x + 0.15, 6.02, 3.7, 0.75, [[(label, {"size": 14})], [(tag or "", {"size": 11, "color": MUTED, "bold": True})]], spacing=1.0)
    x += 4.08

# ---------- 10. Distinction ----------
s = base(10, "One workflow from preparation to outreach.", "The map feeds an inspectable next action.", footer=False)
rows = [("Tool category", "What our design connects"), ("Incident maps", "Geographic context → household outreach queue"), ("Notification systems", "Recipient ordering → visible unresolved contacts"), ("Prevention models", "Treatment exploration → the same incident-area workflow")]
tbl = s.shapes.add_table(len(rows), 2, Inches(0.6), Inches(1.75), Inches(7.6), Inches(2.6)).table
tbl.columns[0].width = Inches(2.6)
tbl.columns[1].width = Inches(5.0)
for r, row in enumerate(rows):
    for c, val in enumerate(row):
        cell = tbl.cell(r, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = PANEL if r else RGBColor(0x3A, 0x31, 0x2C)
        run = cell.text_frame.paragraphs[0].add_run()
        run.text = val
        run.font.name = FONT
        run.font.size = Pt(17)
        run.font.bold = r == 0 or c == 0
        run.font.color.rgb = INK
text(s, 0.6, 4.6, 7.9, 0.5, [[("Prepare before the fire. Prioritize outreach during it.", {"bold": True, "size": 24, "color": INK})]])
text(s, 0.6, 5.2, 7.9, 0.45, [[("Next: a consenting-volunteer pilot and verified call/SMS integration.", {"size": 18, "color": MUTED})]])
text(s, 0.6, 5.85, 7.9, 0.4, [[(F["team"], {"size": 14, "color": MUTED})]])
picture(s, APP / "P3-full.png", 8.55, 1.75, w=4.2)
qr = segno.make(F["repo"], error="m")
qr_path = OUT / "repo-qr.png"
qr.save(str(qr_path), scale=10, border=2, dark="#E5E7E1", light="#1E1A17")
picture(s, qr_path, 8.55, 4.75, w=1.75, border=False)
text(s, 10.45, 4.85, 2.4, 1.4, [[("Code repository", {"bold": True, "size": 14})], [(F["repo"].replace("https://", ""), {"size": 11, "color": MUTED})]], spacing=1.0)
text(s, 0.6, 7.06, 11.4, 0.3, [[("Category comparison, not a claim about named products. Fictional demo contacts; unofficial model outputs. Follow official evacuation instructions.", {"size": 12, "color": MUTED})]])

OUT.mkdir(exist_ok=True)
prs.save(OUT / "WildfireTwinLA-pitch.pptx")
print("saved", OUT / "WildfireTwinLA-pitch.pptx", len(prs.slides), "slides")
